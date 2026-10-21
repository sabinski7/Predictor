"""Generează rapoarte/strategii_verificare.xlsx: strategiile cu numerele celor 10 reguli și cu propunerile, tură cu tură.
Rulare din rădăcina repo-ului: python cercetare/raport_strategii_xlsx.py (necesită openpyxl)."""
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

t = pd.read_csv("top_reguli.csv", dtype=str); t = t[t.iesit.notna()]
lg = pd.read_csv("predictions_log.csv", dtype=str).drop_duplicates("slot", keep="last").set_index("slot")
slots = sorted(t.slot.unique())
rows = []
for s in slots:
    g = t[t.slot == s].sort_values("loc", key=lambda c: c.astype(int))
    rules = g.prezis.astype(int).tolist()
    props = [int(v) for v in lg.loc[s, "propuneri"].split()]
    uniq = list(dict.fromkeys(rules))
    union = list(dict.fromkeys(uniq + props))
    rows.append(dict(slot=s, y=int(g.iesit.iloc[0]), rules=rules, props=props,
                     **{f"c{k}": union[:k] for k in (9, 10, 11, 12)}))
N = len(rows)

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="305496")
B = Font(name=F); BL = Font(name=F, color="0000FF"); BB = Font(name=F, bold=True)
Y = PatternFill("solid", fgColor="FFFF00")
wb = Workbook()

# ---------- Ture ----------
ws = wb.active; ws.title = "Ture"
strat = [  # (cod, nume, câmp, lățime)
    ("A", "Numerele celor 10 reguli", "rules", 10),
    ("B", "Cele 9 propuneri", "props", 9),
    ("C", "Reguli + propuneri până la 9", "c9", 9),
    ("D", "Reguli + propuneri până la 10", "c10", 10),
    ("E", "Reguli + propuneri până la 11", "c11", 11),
    ("F", "Reguli + propuneri până la 12", "c12", 12),
]
hdr = ["Tura (ora Italiei)", "Ieșit"]
col = 3; blocks = {}
for code, nm, key, w in strat:
    blocks[code] = (col, col + w - 1)
    hdr += [f"{code}{i+1}" for i in range(w)]
    col += w
res_start = col
for code, nm, key, w in strat:
    hdr += [f"{code} nr. distincte", f"{code} nimerit", f"{code} pauză"]
ws.append(hdr)
for c in range(1, len(hdr) + 1):
    x = ws.cell(1, c); x.font = H; x.fill = HF; x.alignment = Alignment(horizontal="center", wrap_text=True)
for i, r in enumerate(rows):
    R = i + 2
    vals = [r["slot"], r["y"]]
    for code, nm, key, w in strat:
        v = r[key] + [None] * (w - len(r[key]))
        vals += v
    ws.append(vals)
    for c in range(1, res_start):
        ws.cell(R, c).font = BL if c >= 2 else B
    c = res_start
    for code, *_ in strat:
        a, b = blocks[code]; rng = f"{L(a)}{R}:{L(b)}{R}"
        ws.cell(R, c, "=" + "+".join(f"(COUNTIF({rng},{v})>0)" for v in range(1, 21)))
        ws.cell(R, c + 1, f"=IF(ISNUMBER(MATCH($B{R},{rng},0)),1,0)")
        prev = "0" if R == 2 else f"{L(c+2)}{R-1}"
        ws.cell(R, c + 2, f"=IF({L(c+1)}{R}=1,0,{prev}+1)")
        for k in range(3): ws.cell(R, c + k).font = B
        c += 3
ws.freeze_panes = "C2"
ws.column_dimensions["A"].width = 18
for c in range(2, len(hdr) + 1):
    ws.column_dimensions[L(c)].width = 5 if c < res_start else 11
ws.row_dimensions[1].height = 45
last = N + 1

# ---------- Rezumat ----------
rz = wb.create_sheet("Rezumat", 0)
rz["A1"] = "Verificarea strategiilor pe jurnalul tău"; rz["A1"].font = Font(name=F, bold=True, size=14)
rz["A2"] = f"{N} ture, {rows[0]['slot']} – {rows[-1]['slot']}. Miză: 1 leu pe tură, pe toate numerele strategiei ca un singur bilet."
rz["A2"].font = B
cols = ["Cod", "Strategie", "Numere jucate (medie)", "Ture", "Nimeriri", "Nimerire %", "Așteptat din hazard %",
        "Diferență (pp)", "Cea mai lungă pauză", "Pauze de 8+ ture", "Pauze de 10+ ture", "Cota", "Profit (lei)", "Randament %"]
rz.append([]); rz.append(cols)
for c in range(1, len(cols) + 1):
    x = rz.cell(4, c); x.font = H; x.fill = HF; x.alignment = Alignment(horizontal="center", wrap_text=True)
odds = {"B": 2.11, "C": 2.11, "D": 1.90, "E": 1.73, "F": 1.58}
c = res_start
for j, (code, nm, key, w) in enumerate(strat):
    R = 5 + j
    dc, hc, pc = L(c), L(c + 1), L(c + 2)
    rz.cell(R, 1, code); rz.cell(R, 2, nm)
    rz.cell(R, 3, f"=AVERAGE(Ture!{dc}2:{dc}{last})")
    rz.cell(R, 4, f"=COUNT(Ture!B2:B{last})")
    rz.cell(R, 5, f"=SUM(Ture!{hc}2:{hc}{last})")
    rz.cell(R, 6, f"=E{R}/D{R}")
    rz.cell(R, 7, f"=C{R}/20")
    rz.cell(R, 8, f"=(F{R}-G{R})*100")
    rz.cell(R, 9, f"=MAX(Ture!{pc}2:{pc}{last})")
    rz.cell(R, 10, f"=COUNTIF(Ture!{pc}2:{pc}{last},8)")
    rz.cell(R, 11, f"=COUNTIF(Ture!{pc}2:{pc}{last},10)")
    if code in odds:
        rz.cell(R, 12, odds[code]).font = BL; rz.cell(R, 12).fill = Y
        rz.cell(R, 13, f"=E{R}*L{R}-D{R}")
        rz.cell(R, 14, f"=M{R}/D{R}")
    else:
        rz.cell(R, 12, "variază"); rz.cell(R, 13, "-"); rz.cell(R, 14, "-")
        rz.cell(R, 12).comment = Comment("Regulile dau 6–10 numere diferite de la o tură la alta, deci cota se schimbă la fiecare tură; nu am o cotă pentru fiecare număr de numere.", "Claude")
    for k in range(1, 15):
        if rz.cell(R, k).font != BL: rz.cell(R, k).font = B
    rz.cell(R, 3).number_format = "0.0"; rz.cell(R, 6).number_format = "0.0%"; rz.cell(R, 7).number_format = "0.0%"
    rz.cell(R, 8).number_format = "+0.0;-0.0;0.0"; rz.cell(R, 13).number_format = "+0.00;-0.00;0.00"; rz.cell(R, 14).number_format = "+0.0%;-0.0%;0.0%"
    c += 3
rz.cell(5, 12).comment = None

# random baseline (simulare)
rng = np.random.default_rng(0); y = np.array([r["y"] for r in rows])
R = 12
rz.cell(R - 1, 1, "Comparație: numere alese la întâmplare (10.000 de simulări pe aceleași rezultate)").font = BB
rz.append(["Numere", "Nimerire % medie", "Cea mai lungă pauză (mediană)", "Cea mai lungă pauză: 90% din cazuri între", ""])
for c_ in range(1, 5):
    x = rz.cell(R, c_); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True)
for k in (8, 9, 10, 11, 12):
    mx = []; hr = []
    for _ in range(10000):
        h = rng.random((N, 20)).argsort(1)[:, :k] + 1 == y[:, None]
        h = h.any(1); hr.append(h.mean())
        r_ = 0; m = 0
        for v in h:
            r_ = 0 if v else r_ + 1; m = max(m, r_)
        mx.append(m)
    rz.append([k, float(np.mean(hr)), int(np.median(mx)), f"{int(np.percentile(mx,5))} – {int(np.percentile(mx,95))}"])
    rr = rz.max_row
    for c_ in range(1, 5): rz.cell(rr, c_).font = B
    rz.cell(rr, 2).number_format = "0.0%"

n0 = rz.max_row + 2
notes = [
    "Cum se citește",
    "• Nimerire %: în câte ture numărul ieșit a fost printre numerele strategiei.",
    "• Așteptat din hazard: numere jucate × 5% (o extragere 1–20 perfect aleatoare).",
    "• Pauză: câte ture la rând fără nimerire. „Pauze de 8+” numără câte serii de ratări au ajuns la 8 sau mai mult.",
    "• Cota (celulele galbene, albastre) e cea dată de tine: 9 numere 2,11 · 10 numere 1,90 · 11 numere 1,73 · 12 numere 1,58. O poți modifica; profitul se recalculează.",
    "• Profit: miză 1 leu pe tură (biletul cu toate numerele strategiei), câștig = cota × 1 leu la nimerire.",
    "• Strategiile C–F: întâi numerele regulilor (în ordinea din top, fără dubluri), apoi propunerile notebook-ului, până se adună numărul de numere.",
    "  F (12) are uneori mai puțin de 12 numere, când regulile și propunerile se suprapun mult.",
    "• Top-ul regulilor la fiecare tură e calculat doar din turele anterioare (exact ce ai fi primit atunci). Sursa: top_reguli.csv și predictions_log.csv.",
    "• Foaia „Ture” are numerele fiecărei strategii (albastru = date) și formulele de nimerire și pauză.",
]
for i, s_ in enumerate(notes):
    rz.cell(n0 + i, 1, s_).font = BB if i == 0 else B
widths = [10, 32, 12, 8, 10, 11, 12, 11, 12, 11, 11, 9, 11, 11]
for i, w_ in enumerate(widths, 1): rz.column_dimensions[L(i)].width = w_
rz.row_dimensions[4].height = 45; rz.row_dimensions[R].height = 45
wb.calculation.fullCalcOnLoad = True
wb.save("rapoarte/strategii_verificare.xlsx")
print(N, "ture")
